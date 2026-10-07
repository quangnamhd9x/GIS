"""
Dựng dữ liệu bản đồ KD Green Farm từ file gốc của nông trường.

Đầu vào:
  --excel  Định danh vườn cây (.xlsx, sheet "Dinhdanh_KDGF" + "Data")
  --lots   PDF "Bản đồ thể hiện cụm sản xuất A,B và cập nhật tên lô" (VN-2000, 1:10.000)
  --plan   PDF "Bản đồ dự kiến bàn giao KD 3 giai đoạn năm 2026" (tùy chọn)
Đầu ra (thư mục --out):
  lots.geojson        mỗi lô một Feature (Polygon nếu ranh giới PDF khớp diện tích Excel, ngược lại Point)
  plan-2026.geojson   các vùng kế hoạch thu hồi / hiện trạng năm 2026
  data-report.json    tóm tắt + cảnh báo dữ liệu để rà soát

Chạy:  pip install openpyxl pymupdf shapely pyproj
       python tools/build_farm_data.py --excel ... --lots ... --plan ... --out react-farm-map/public/data --js data/lots.js
"""
import argparse
import collections
import datetime as dt
import json
import re
from pathlib import Path

import openpyxl
import pymupdf
from pyproj import Transformer
from shapely.geometry import Point, Polygon, mapping
from shapely.ops import transform as shp_transform, unary_union

from pdf_paths import fill_paths

# VN-2000, múi chiếu 3°, kinh tuyến trục Đắk Lắk 108°30' -> WGS84
VN2000_DAKLAK = (
    "+proj=tmerc +lat_0=0 +lon_0=108.5 +k=0.9999 +x_0=500000 +y_0=0 +ellps=WGS84 "
    "+towgs84=-191.90441429,-39.30318279,-111.45032835,-0.00928836,0.01975479,-0.00427372,0.252906278 +units=m +no_defs"
)
TO_WGS84 = Transformer.from_proj(VN2000_DAKLAK, "EPSG:4326", always_xy=True)

# Hiệu chỉnh lưới tọa độ đọc từ nhãn khung bản đồ (điểm PDF -> mét VN-2000)
LOTS_GRID = dict(x0=276.55, e0=491000, y0=211.6, n0=1401000, pt_per_km_x=305.0, pt_per_km_y=304.95)
PLAN_GRID = dict(x0=274.6, e0=494000, y0=438.3, n0=1400000, pt_per_km_x=283.45, pt_per_km_y=283.4)

# Polygon PDF chỉ được dùng khi diện tích lệch Excel trong khoảng này
AREA_RATIO_OK = (0.8, 1.25)


def pdf_to_wgs(grid):
    def f(x, y):
        e = grid["e0"] + (x - grid["x0"]) / grid["pt_per_km_x"] * 1000
        n = grid["n0"] - (y - grid["y0"]) / grid["pt_per_km_y"] * 1000
        return TO_WGS84.transform(e, n)
    return f


def pdf_area_ha(geom, grid):
    return geom.area / (grid["pt_per_km_x"] * grid["pt_per_km_y"]) * 100  # km² -> ha


def round_geom(geom):
    """GeoJSON với tọa độ làm tròn 6 chữ số (~0,1 m)."""
    def r(c):
        if isinstance(c, (list, tuple)) and c and isinstance(c[0], (int, float)):
            return [round(c[0], 6), round(c[1], 6)]
        return [r(x) for x in c]
    g = mapping(geom)
    return {"type": g["type"], "coordinates": r(g["coordinates"])}


def num(v):
    return v if isinstance(v, (int, float)) else None


def iso(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else None


def clean_str(v):
    return str(v).strip() if v not in (None, "") else None


# ---------------------------------------------------------------- Excel
def read_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = collections.defaultdict(list)
    for r in wb["Dinhdanh_KDGF"].iter_rows(min_row=3, values_only=True):
        if r[0] and str(r[0]).startswith("Cụm ") and isinstance(r[4], int):
            rows[r[4]].append(r)

    status_rows = collections.defaultdict(list)
    for r in wb["Data"].iter_rows(min_row=3, values_only=True):
        if r[1] and str(r[1]).startswith("Cụm ") and isinstance(r[5], int):
            status_rows[r[5]].append(r)

    lots = {}
    for lo, rs in rows.items():
        first = rs[0]
        s = lambda i: round(sum(num(r[i]) or 0 for r in rs), 4)
        tai_canh = []
        for dot, (ci, gi, ai, ki) in ((1, (13, 14, 15, 16)), (2, (17, 18, 19, 20))):
            area = s(ai)
            if area:
                tai_canh.append({
                    "dot": dot,
                    "ngay": min((iso(r[ci]) for r in rs if iso(r[ci])), default=None),
                    "giong": next((clean_str(r[gi]) for r in rs if clean_str(r[gi])), None),
                    "dienTich": round(area, 2),
                    "soCay": int(s(ki)),
                })
        dates = sorted(iso(r[7]) for r in rs if iso(r[7]))
        lots[lo] = {
            "lo": lo,
            "loCu": [str(r[3]).strip() for r in rs],
            "cum": first[0].replace("Cụm ", ""),
            "doi": first[2],
            "nongTruong": first[1],
            "giong": next((clean_str(r[8]) for r in rs if clean_str(r[8])), None),
            "ngayTrong": dates[0] if dates else None,
            "matDo": num(first[5]),
            "dienTichBanDau": round(s(9), 2),
            "soCayBanDau": int(s(10)),
            "taiCanh": tai_canh,
            "dienTichHienHuu": round(s(21), 2),
            "soCayHienHuu": int(s(22)),
        }
        # Trạng thái & người phụ trách: dòng có số thứ tự trồng lớn nhất trong sheet Data
        srs = status_rows.get(lo, [])
        if srs:
            top = max(num(r[6]) or 0 for r in srs)
            latest = [r for r in srs if (num(r[6]) or 0) == top]
            st = {r[7] for r in latest}
            lots[lo]["trangThai"] = "Đã trồng" if "Đã trồng" in st else ("Chưa trồng" if "Chưa trồng" in st else "Đã hủy")
            lots[lo]["phuTrach"] = clean_str(latest[0][15])
            lots[lo]["maLo"] = clean_str(latest[0][0])
            note = clean_str(latest[0][17])
            if note:
                lots[lo]["ghiChu"] = note
        if not lots[lo]["giong"] and tai_canh:
            lots[lo]["giong"] = tai_canh[-1]["giong"]
    return lots


# ---------------------------------------------------------------- PDF lô
def read_lot_map(path):
    doc = pymupdf.open(path)
    page = doc[0]
    out, _ = fill_paths(path)
    polys = []
    for fill, rings in out:
        if fill in ("/P15", "/P18"):  # pattern tô lô Cụm A / Cụm B
            g = unary_union([Polygon(r).buffer(0) for r in rings if len(r) >= 3])
            if not g.is_empty:
                polys.append(g)

    # Nhãn "Lô 54" (font VNI hiển thị thành "Loâ")
    lines = collections.defaultdict(list)
    for w in page.get_text("words"):
        lines[(w[5], w[6])].append(w)
    labels = collections.defaultdict(list)
    for ws in lines.values():
        ws.sort(key=lambda w: w[0])
        for i, w in enumerate(ws):
            if w[4] == "Loâ" and i + 1 < len(ws) and ws[i + 1][4].isdigit():
                labels[int(ws[i + 1][4])].append(((w[0] + ws[i + 1][2]) / 2, (w[1] + w[3]) / 2))
    return polys, labels


# ---------------------------------------------------------------- PDF kế hoạch 2026
PLAN_CLASSES = {
    (0.81, 1.0, 0.5): ("chuoi", "Diện tích đã trồng chuối", None),
    (1.0, 0.82, 0.63): ("caoSuDucLoc", "Cao su Đức Lộc", None),
    (1.0, 0.86, 0.44): ("caoSuLK", "Cao su giao khoán cho LK", None),
    (0.0, 1.0, 1.0): ("gd1", "Dự kiến thu hồi GĐ1", 30.75),
    (1.0, 1.0, 0.0): ("gd2", "Dự kiến thu hồi GĐ2", 40.8),
    (1.0, 0.57, 0.57): ("gd3", "Dự kiến thu hồi GĐ3", 69.76),
    (1.0, 0.5, 0.75): ("traiHeo", "Trang trại nuôi heo", None),
    (0.82, 0.41, 0.0): ("caoToc", "Cao tốc BMT - KH", None),
}


def read_plan_map(path):
    page = pymupdf.open(path)[0]
    to_wgs = pdf_to_wgs(PLAN_GRID)
    seen, feats = set(), []
    for d in page.get_drawings():
        f = d.get("fill")
        if f is None:
            continue
        key = tuple(round(v, 2) for v in f)
        if key not in PLAN_CLASSES:
            continue
        r = d["rect"]
        if (r.x0 < 520 and r.y0 > 2050) or r.width > 1500:  # bỏ ô chú giải & khung
            continue
        sig = (key, round(r.x0), round(r.y0), round(r.x1), round(r.y1))
        if sig in seen:
            continue
        seen.add(sig)
        pts = []
        for it in d["items"]:
            if it[0] == "l":
                pts += [it[1], it[2]]
            elif it[0] == "c":
                pts += [it[1], it[4]]
            elif it[0] == "re":
                q = it[1]
                pts += [q.tl, q.tr, q.br, q.bl]
        if len(pts) < 3:
            continue
        geom = Polygon([(p.x, p.y) for p in pts]).buffer(0)
        if geom.is_empty or geom.area < 4:
            continue
        cls, name, legend_ha = PLAN_CLASSES[key]
        feats.append({
            "type": "Feature",
            "geometry": round_geom(shp_transform(to_wgs, geom).simplify(0.000003)),
            "properties": {"loai": cls, "ten": name, "dienTichChuGiai": legend_ha},
        })
    return {"type": "FeatureCollection", "features": feats}


# ---------------------------------------------------------------- Ghép dữ liệu
def build(excel, lots_pdf, plan_pdf, out_dir, js_path=None):
    lots = read_excel(excel)
    polys, labels = read_lot_map(lots_pdf)
    to_wgs = pdf_to_wgs(LOTS_GRID)
    issues = []

    # Đếm polygon được bao nhiêu nhãn lô "chiếm" để loại polygon bị gộp nhiều lô
    owner = collections.defaultdict(set)
    for lo, pts in labels.items():
        for x, y in pts:
            for i, g in enumerate(polys):
                if g.contains(Point(x, y)):
                    owner[i].add(lo)

    features = []
    for lo in sorted(lots):
        p = lots[lo]
        pts = labels.get(lo, [])
        geom, source = None, "none"
        if pts:
            best = None
            for x, y in pts:
                for i, g in enumerate(polys):
                    if g.contains(Point(x, y)) and len(owner[i]) == 1:
                        ratio = pdf_area_ha(g, LOTS_GRID) / p["dienTichBanDau"] if p["dienTichBanDau"] else 0
                        if AREA_RATIO_OK[0] <= ratio <= AREA_RATIO_OK[1] and (best is None or abs(ratio - 1) < abs(best[1] - 1)):
                            best = (g, ratio)
            cx = sum(x for x, _ in pts) / len(pts)
            cy = sum(y for _, y in pts) / len(pts)
            if best:
                g = best[0]
                geom = shp_transform(to_wgs, g).simplify(0.000003)
                c = g.representative_point()
                p["viTri"] = [round(v, 6) for v in to_wgs(c.x, c.y)[::-1]]
                p["dienTichBanDo"] = round(pdf_area_ha(g, LOTS_GRID), 2)
                source = "pdf"
            else:
                p["viTri"] = [round(v, 6) for v in to_wgs(cx, cy)[::-1]]
                geom = Point(to_wgs(cx, cy))
                source = "nhan"
        else:
            issues.append({"lo": lo, "muc": "vi-tri", "noiDung": f"Lô {lo} có trong Excel nhưng không tìm thấy trên bản đồ PDF"})
        p["nguonRanhGioi"] = source
        if not p["giong"]:
            issues.append({"lo": lo, "muc": "excel", "noiDung": f"Lô {lo} chưa ghi loại giống"})
        if len(p["loCu"]) > 1:
            issues.append({"lo": lo, "muc": "gop-lo", "noiDung": f"Lô {lo} gộp từ {len(p['loCu'])} lô cũ: {', '.join(p['loCu'])}"})
        if geom is not None:
            features.append({"type": "Feature", "geometry": round_geom(geom), "properties": p})

    numbers = sorted(lots)
    missing = [n for n in range(numbers[0], numbers[-1] + 1) if n not in lots]
    if missing:
        issues.append({"lo": None, "muc": "thieu-so", "noiDung": "Không có lô số: " + ", ".join(map(str, missing))})
    on_map_only = sorted(set(labels) - set(lots))
    if on_map_only:
        issues.append({"lo": None, "muc": "vi-tri", "noiDung": "Có trên bản đồ nhưng không có trong Excel: " + ", ".join(map(str, on_map_only))})

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    (out / "lots.geojson").write_text(json.dumps({"type": "FeatureCollection", "features": features}, ensure_ascii=False), encoding="utf-8")
    if plan_pdf:
        (out / "plan-2026.geojson").write_text(json.dumps(read_plan_map(plan_pdf), ensure_ascii=False), encoding="utf-8")

    src = collections.Counter(f["properties"]["nguonRanhGioi"] for f in features)
    report = {
        "taoLuc": dt.datetime.now().isoformat(timespec="minutes"),
        "nguon": {"excel": Path(excel).name, "banDoLo": Path(lots_pdf).name, "banDoKeHoach": Path(plan_pdf).name if plan_pdf else None},
        "tongLo": len(lots),
        "coRanhGioi": src.get("pdf", 0),
        "chiCoViTri": src.get("nhan", 0),
        "khongCoViTri": len(lots) - len(features),
        "tongDienTichHienHuu": round(sum(p["dienTichHienHuu"] for p in lots.values()), 2),
        "tongSoCayHienHuu": sum(p["soCayHienHuu"] for p in lots.values()),
        "canhBao": issues,
    }
    (out / "data-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if js_path:
        # Bản rút gọn cho trang HTML tĩnh (mở trực tiếp bằng file:// không fetch được JSON)
        lite = [{k: lots[n].get(k) for k in ("lo", "loCu", "cum", "doi", "giong", "dienTichHienHuu", "soCayHienHuu", "trangThai", "viTri")}
                for n in sorted(lots) if lots[n].get("viTri")]
        Path(js_path).write_text("// Tự sinh bởi tools/build_farm_data.py - không sửa tay\nwindow.KDGF_LOTS = "
                                 + json.dumps(lite, ensure_ascii=False) + ";\n", encoding="utf-8")
    return report


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--excel", required=True)
    ap.add_argument("--lots", required=True)
    ap.add_argument("--plan")
    ap.add_argument("--out", default="react-farm-map/public/data")
    ap.add_argument("--js", help="(tùy chọn) xuất thêm file JS cho trang HTML tĩnh, VD data/lots.js")
    a = ap.parse_args()
    r = build(a.excel, a.lots, a.plan, a.out, a.js)
    print(json.dumps({k: v for k, v in r.items() if k != "canhBao"}, ensure_ascii=False, indent=2))
    print(f"{len(r['canhBao'])} cảnh báo -> data-report.json")
